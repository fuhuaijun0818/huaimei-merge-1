from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

fonts = ["Noto Sans CJK SC", "STSong", "Songti SC", "Arial Unicode MS", "Hiragino Sans GB", "PingFang SC", "SimSun", "Microsoft YaHei"]
doc = Document()
for font in fonts:
    p = doc.add_paragraph()
    r = p.add_run(f"{font}: 华美一期售前智能体产品设计方案")
    r.font.name = font
    r.font.size = Pt(18)
    r._element.rPr.rFonts.set(qn("w:ascii"), font)
    r._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    r._element.rPr.rFonts.set(qn("w:eastAsia"), font)
doc.save("/private/tmp/huamei-font-test.docx")
