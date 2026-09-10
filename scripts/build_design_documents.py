from __future__ import annotations

from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT.parent / "交付文档"
OUT.mkdir(parents=True, exist_ok=True)

FONT = "Arial"
CJK_FONT = "PingFang SC"
INK = "252A2E"
MUTED = "6B7075"
BRAND = "7A2742"
BRAND_LIGHT = "F8EFF2"
BRAND_MID = "DDB4BF"
PALE = "F4F6F8"
BLUE_LIGHT = "E8EEF5"
GREEN = "347A5A"
AMBER = "8A6200"
RED = "9B1C1C"
WHITE = "FFFFFF"
TABLE_WIDTH = 9360


def set_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for col in list(grid):
        grid.remove(col)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_borders(table, color="D9DDE1", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def format_cell_text(cell, text, size=9.5, color=INK, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(str(text))
    set_font(run, size=size, color=color, bold=bold)


def configure_document(doc, title, subtitle):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.82)
    sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.35)
    sec.footer_distance = Inches(0.34)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18
    for style_name, size, color, before, after in (
        ("Heading 1", 16, BRAND, 18, 8),
        ("Heading 2", 12.5, BRAND, 12, 6),
        ("Heading 3", 11, "4F2633", 8, 4),
    ):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.18
    header = sec.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hp.paragraph_format.space_after = Pt(0)
    hrun = hp.add_run(f"上海华美医疗美容医院｜{subtitle}")
    set_font(hrun, size=8.5, color=MUTED)
    footer = sec.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    frun = fp.add_run("华美一期｜内部评审稿｜第 ")
    set_font(frun, size=8.5, color=MUTED)
    add_field(frun, "PAGE")
    frun2 = fp.add_run(" 页")
    set_font(frun2, size=8.5, color=MUTED)
    props = doc.core_properties
    props.title = title
    props.subject = "华美一期售前智能体"
    props.author = "项目组"
    props.comments = "内部评审稿"


def add_field(run, instruction):
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r.append(begin); r.append(instr); r.append(separate); r.append(text); r.append(end)


def add_cover(doc, title, subtitle, note):
    doc.add_paragraph().paragraph_format.space_after = Pt(30)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(10)
    r = kicker.add_run("时代之美｜客户全生命周期智能经营平台")
    set_font(r, size=10.5, color=BRAND, bold=True)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_after = Pt(8)
    r = t.add_run(title)
    set_font(r, size=26, color=BRAND, bold=True)
    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.paragraph_format.space_after = Pt(26)
    r = st.add_run(subtitle)
    set_font(r, size=13, color=MUTED)
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_geometry(table, [1900, 7460])
    set_borders(table, color="E3D7DB")
    metadata = [
        ("客户", "上海华美医疗美容医院"),
        ("版本", "V1.0｜内部评审稿"),
        ("日期", "2026年9月3日"),
        ("文件定位", note),
    ]
    for row, (label, value) in zip(table.rows, metadata):
        shade(row.cells[0], BRAND_LIGHT)
        format_cell_text(row.cells[0], label, size=10, color=BRAND, bold=True)
        format_cell_text(row.cells[1], value, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(70)
    statement = doc.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.CENTER
    statement.paragraph_format.space_after = Pt(0)
    r = statement.add_run("本文件用于一期范围、设计原则与实施准入评审；不构成医疗建议、价格承诺或外部系统执行承诺。")
    set_font(r, size=9.5, color=MUTED, italic=True)
    doc.add_page_break()


def add_para(doc, text, bold_prefix=None, color=INK, size=10.5, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.18
    if keep:
        p.paragraph_format.keep_with_next = True
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, size=size, color=color, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r, size=size, color=color)
    else:
        r = p.add_run(text)
        set_font(r, size=size, color=color)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_font(r, size=10.5, color=INK)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_font(r, size=10.5, color=INK)


def add_section_title(doc, heading, lead=None):
    doc.add_heading(heading, level=1)
    if lead:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.18
        r = p.add_run(lead)
        set_font(r, size=10.5, color="45494D", italic=True)


def add_subtitle(doc, heading, text=None):
    doc.add_heading(heading, level=2)
    if text:
        add_para(doc, text)


def add_callout(doc, label, text, tone="brand"):
    color = BRAND_LIGHT if tone == "brand" else ("F5F7FA" if tone == "neutral" else "FFF7E6")
    edge = BRAND if tone == "brand" else ("52677F" if tone == "neutral" else AMBER)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, [9360])
    set_borders(table, color=edge, size="8")
    cell = table.cell(0, 0)
    shade(cell, color)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(label + "  ")
    set_font(r, size=10, color=edge, bold=True)
    r = p.add_run(text)
    set_font(r, size=10, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows, widths, header_fill=PALE, font_size=9.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_geometry(table, widths)
    set_borders(table)
    for cell, header in zip(table.rows[0].cells, headers):
        shade(cell, header_fill)
        format_cell_text(cell, header, size=9.2, color=BRAND, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            format_cell_text(cell, value, size=font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_process(doc, steps):
    table = doc.add_table(rows=1, cols=len(steps))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    widths = [TABLE_WIDTH // len(steps)] * len(steps)
    widths[-1] += TABLE_WIDTH - sum(widths)
    set_table_geometry(table, widths)
    set_borders(table, color="E4CCD2")
    for idx, (title, detail) in enumerate(steps):
        cell = table.cell(0, idx)
        shade(cell, BRAND_LIGHT if idx % 2 == 0 else "FFF9FA")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{idx + 1}. {title}")
        set_font(r, size=9.2, color=BRAND, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.05
        r = p2.add_run(detail)
        set_font(r, size=8.5, color=MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def product_document():
    doc = Document()
    configure_document(doc, "华美一期售前智能体产品设计方案", "产品设计方案")
    add_cover(doc, "华美一期售前智能体\n产品设计方案", "以可信知识、自然对话和人工协同降低售前人工投入", "产品范围、运营闭环与验收设计")

    add_section_title(doc, "1. 产品定位与目标", "一期产品面向上海华美医疗美容医院的售前咨询场景，先完成可控、可追溯、可运营的独立系统闭环。")
    add_subtitle(doc, "1.1 北极星目标")
    add_para(doc, "在咨询量口径一致、服务水平不降低、医疗与合规边界不放宽的前提下，减少售前人员在首轮接待、知识查询、基础解释、重复澄清和留资整理上的实际付出时间。具体降幅须以上线前后同口径基线数据确认；30%仅作为当前建议目标，不作为未经验证的承诺。")
    add_table(doc, ["目标类别", "一期目标", "衡量方式"], [
        ("产品目标", "建立具备中等水平售前人员业务能力的智能体，能识别意图、承接多轮、调用企业知识、自然表达并安全转人工。", "核心意图、槽位、多轮状态、事实证据与人工盲评。"),
        ("运营目标", "在相同服务标准下减少人工单客付出分钟数，并形成可复盘的 Bad Case 与知识更新闭环。", "连续四周人工分钟、首响、转人工、满意度、问题闭环时效。"),
        ("开发目标", "形成前台对话、管理后台、知识治理、评测、审计、部署与回滚能力；支持后续多渠道扩展。", "功能验收、回归评测、20并发长稳、演练与发布证据。"),
    ], [1600, 4700, 3060])
    add_callout(doc, "范围边界", "一期以网页咨询和内部运营后台为真实交付范围。企业微信、CRM、外部预约、支付等保持接口预留，不以模拟成功替代真实对接。", "neutral")

    add_section_title(doc, "2. 用户、角色与关键任务")
    add_table(doc, ["用户", "典型任务", "一期产品价值"], [
        ("终端客户", "了解项目、医生、产品、公开资料、到院流程；表达顾虑并提出预约意向。", "获得及时、自然且有依据的初步咨询；必要时进入人工服务。"),
        ("售前顾问/前台", "接管高风险、复杂或预约相关会话；核验排班；登记工时与处理结果。", "减少重复解释，接管时获得上下文摘要、证据和待办字段。"),
        ("内容专员/科室主管", "上传或维护项目、医生、产品、价格、活动及公开素材。", "以版本、标签、有效期和权限维护内容，不直接影响对客回答。"),
        ("业务、医务、合规审核人", "审核业务准确性、医疗表述、宣传与隐私边界。", "三审留痕，明确资料是否可对客使用。"),
        ("发布与索引管理员", "创建发布包、查看建图/评测结果、激活或回滚版本。", "保证客户只检索已批准且可追溯的知识版本。"),
    ], [1650, 4000, 3710])

    add_section_title(doc, "3. 信息架构与产品范围")
    add_table(doc, ["端", "模块", "核心能力"], [
        ("客户前台", "AI获客顾问", "首次接待、多轮咨询、项目/医生/产品资料、媒体预览、预约意向、人工接管、回答赞踩。"),
        ("运营后台", "会话与接待", "会话列表、接管队列、预约提醒、排班核验、人工回复、工时与处理结果。"),
        ("运营后台", "知识与合规", "资料录入、版本、审核、发布、对客范围、有效期、GraphRAG版本与回滚。"),
        ("运营后台", "产品与医生", "可维护的产品、价格字段、医生资料、公开素材与来源证据。"),
        ("运营后台", "组织与权限", "用户、角色、科室范围、操作审计与职责分离。"),
        ("运营后台", "质量与分析", "赞踩 Bad Case、评测、召回/证据分析、运营指标和告警。"),
    ], [1450, 2100, 5810])

    add_section_title(doc, "4. 客户咨询旅程与对话体验")
    add_process(doc, [
        ("接待", "固定欢迎语\n说明可提供的服务"),
        ("理解", "识别项目、部位、方式、情绪与问题"),
        ("回答", "按路由获取事实\n短句分段表达"),
        ("澄清", "仅补问最关键\n1—2项信息"),
        ("转化", "留资、预约意向\n或人工接管"),
        ("闭环", "赞踩反馈、Bad Case\n持续优化"),
    ])
    add_subtitle(doc, "4.1 对话体验原则")
    add_bullets(doc, [
        "先回应用户已表达的诉求，再给出相关信息；不得忽略已确认的项目、方式或偏好后重复泛问。",
        "采用短句、分段、自然称呼和适度情绪回应；编号内容保持同一消息块，避免机械逐条发送。",
        "可在合适场景给予非医疗、非夸大式的肯定与安抚；情绪表达不得改变医学事实、价格或权限结论。",
        "医生、项目、价格、活动、院区和公开素材等企业事实必须引用内部已发布知识；通用科普可使用模型通识，但须保持医疗边界。",
        "客户可点击回答引用的图片或PDF预览公开资料；不可向客户展示内部审核材料、客户案例或受限媒体。",
    ])
    add_subtitle(doc, "4.2 首次接待与意图承接")
    add_para(doc, "当用户发送“你好”等问候，智能体以固定身份完成介绍，说明可咨询医美项目、预约面诊、会员权益及活动通知，并告知上海市区首次到院可协助安排网约车。随后仅询问用户感兴趣的项目，不进入泛化部位追问。")
    add_table(doc, ["典型输入", "应识别的核心意图", "期望行为"], [
        ("我想丰胸", "项目咨询：胸部整形", "介绍可选方向，并就方式、自然度、恢复或预算做最小必要澄清。"),
        ("自体脂肪方式", "待选择槽位：丰胸方式=自体脂肪", "确认已选方式，直接介绍其原理与下一步可问主题，不回退到泛化项目推荐。"),
        ("李健院长有哪些项目", "医生咨询+项目范围", "从医生知识和项目关联中回答；排班、职称等时效事实明确核验。"),
        ("面膜", "产品目录明细", "识别为上轮产品/明星产品追问，介绍对应产品并展示公开图片资料。"),
    ], [2000, 2600, 4760])

    add_section_title(doc, "5. 售前智能体能力设计")
    add_table(doc, ["能力", "产品规则", "用户可感知结果"], [
        ("意图与槽位", "识别项目、部位、方式、关注点、预算、时间、医生、产品及预约诉求；用通用待选择槽位机制承接简短选择。", "不会因“第二种”“自体脂肪方式”等短句丢失上下文。"),
        ("对话状态", "保存当前目标、已确认信息、待补字段、已完成步骤和话题版本；纠正或换题时覆盖旧状态。", "减少重复提问，客户插话后仍能继续原任务。"),
        ("动态路由", "区分通识、知识库、结构化数据、GraphRAG、媒体、人工服务和受控写操作。", "不同问题走合适链路，不把全部问题都交给同一模型。"),
        ("证据与时效", "关键事实绑定来源、版本、生效/失效时间和适用范围；无法核验的动态事实不作确定承诺。", "客户可查看公开引用；过期价格或排班不会被当作当前事实。"),
        ("多模态资料", "按权限展示已批准的医生/项目海报、图片和PDF；媒体与文字证据同版本治理。", "可直接预览相关公开资料。"),
        ("人工协同", "投诉、医疗高风险、复杂审批、预约确认等携带摘要、证据和已采集字段进入人工队列。", "AI不越权承诺，人工无需从头了解客户。"),
    ], [1600, 4700, 3060])

    add_section_title(doc, "6. 知识、产品与内容运营")
    add_subtitle(doc, "6.1 知识分层")
    add_table(doc, ["知识层", "示例", "对客规则"], [
        ("通用科普", "常见术式原理、恢复期一般注意事项", "仅用于非个体化教育；提示需由专业医生面诊确认。"),
        ("企业公开知识", "项目、医生介绍、院区、公开产品资料、服务流程", "必须来自已发布知识包并携带来源。"),
        ("结构化业务数据", "产品当前价格、活动、排班、名额、预约状态", "优先从权威结构化源读取；未接入时明确需人工核验。"),
        ("受限内部知识", "运营汇总、审核意见、客户案例、聊天记录", "不进入客户检索，不作为普通RAG或GraphRAG上下文。"),
    ], [1650, 3850, 3960])
    add_subtitle(doc, "6.2 后台内容更新流程")
    add_process(doc, [
        ("录入", "上传或修订\n填写来源/有效期"),
        ("解析", "提取文本、标签、\n重复与风险提示"),
        ("审核", "科室确认＋业务、\n医务、合规三审"),
        ("建图", "候选索引、差异\n与回归评测"),
        ("发布", "灰度或正式激活\n旧版本可回滚"),
    ])
    add_callout(doc, "时效性规则", "价格、活动、名额、排班和医生出诊必须有权威来源及生效/失效时间；到期自动停止对客检索。此类数据优先结构化管理，不以GraphRAG作为唯一事实源。", "caution")

    add_section_title(doc, "7. 管理后台与运营闭环")
    add_table(doc, ["工作区", "关键操作", "控制要求"], [
        ("会话与接待", "查看会话、认领人工接管、确认预约、登记工时和处理结论。", "未认领者不能处理；AI接管后静默；联系方式显式查看与审计。"),
        ("预约提醒", "接收预约意向、领取事项、核验内部排班、确认或退回补充。", "意向不等于预约成功；必须选择有效排班并使用幂等键。"),
        ("产品与医生", "维护产品字段、当前价格、有效期、销售状态、医生资料及公开素材。", "价格只是产品字段；原始证据保留为内部依据。"),
        ("知识与合规", "录入、审核、发布、下架、版本比对和GraphRAG绑定。", "资料变更产生新版本；无三审不得对客发布。"),
        ("质量与分析", "查看赞踩、Bad Case、检索证据、评测与运营指标。", "踩评进入优化队列；禁止把客户隐私直接写入开放评测集。"),
    ], [1700, 3900, 3860])

    add_section_title(doc, "8. 安全、医疗与合规体验")
    add_bullets(doc, [
        "智能体不得诊断、替代面诊、承诺疗效、承诺价格、虚构排班或宣称已完成未接入的业务操作。",
        "敏感医疗、投诉、隐私、紧急风险、复杂会员审核或无法证实的信息，进入明确的不确定性说明或人工服务。",
        "长期偏好、联系方式和主动触达均需用户授权；用户应可查看、更正、删除或关闭授权范围内的数据。",
        "客户资料、图片和受限媒体不进入通用向量检索；日志、引用、评测和Bad Case均需脱敏与权限控制。",
    ])
    add_callout(doc, "对客声明", "页面固定展示“项目信息仅供参考，具体适用性请咨询专业医师”；对涉及医疗决定的问题，回答中提供面诊/人工咨询路径。", "brand")

    add_section_title(doc, "9. 评测与验收设计")
    add_table(doc, ["维度", "生产目标", "主要验证集"], [
        ("理解与状态", "核心意图Top-1≥95%；槽位F1≥95%；多轮状态准确率≥95%；用户纠正采纳率≥98%。", "固定回归集、长对话集、真实问题集。"),
        ("对话质量", "无必要反问率≤5%；拟人度人工评分≥4.3/5。", "人工盲评、赞踩Bad Case与风格回归集。"),
        ("事实与证据", "有依据事实准确率≥95%；引用一致率≥98%；重要声明证据覆盖率≥95%；无依据强答≤2%。", "知识事实集、来源一致性检查、对抗集。"),
        ("安全与任务", "高风险写操作确认覆盖率100%；非授权/非预期操作0；核心任务成功率≥95%。", "对抗安全集、工具执行集、人工接管演练。"),
        ("性能", "简单问答首字P95≤2秒；GraphRAG P95≤2秒；20并发30分钟完整链路稳定。", "压测、长稳、故障注入与监控原始数据。"),
    ], [1600, 4950, 2810], font_size=9.1)

    add_section_title(doc, "10. 版本节奏与上线门禁")
    add_table(doc, ["阶段", "周次", "重点工作", "里程碑/出口"], [
        ("S0 立项与基线", "W1–W2", "目标/范围确认、业务基线、渠道与接口盘点、数据安全边界。", "M0：立项与方案评审通过"),
        ("S1 产品与知识设计", "W2–W3", "场景、流程、能力模型、知识分层、合规规则、评测集设计。", "M1：需求与知识规范冻结"),
        ("S2 开发与知识生产", "W3–W6", "智能体、RAG、GraphRAG准备、知识审核、管理后台与监控。", "M2：集成环境MVP可用"),
        ("S3 联调与验证", "W6–W8", "后台、场景测试、安全/合规测试、压力与降级验证。", "M3：试点准入评审通过"),
        ("S4 灰度试点", "W8–W9", "小流量、人工兜底、日报、Bad Case修复、阈值调整。", "M4：扩大灰度决策"),
        ("S5 验收与交接", "W9–W10", "业务对比、技术验收、培训、运维交接、遗留项计划。", "M5：一期验收/有条件验收"),
    ], [1450, 1000, 4500, 2410], font_size=8.9)
    add_callout(doc, "上线门禁", "正式对客前必须具备：已生效的三审知识版本、正式GraphRAG索引或可审计降级策略、真实模型与20并发证据、身份权限、监控告警、备份恢复、人工接管演练和项目委员会Go决策。", "caution")

    add_section_title(doc, "11. 当前状态与后续决策")
    add_para(doc, "当前系统可在本地开发环境完成对话、后台管理、已审核知识检索、媒体引用、预约意向和质量反馈等验证。正式GraphRAG、真实模型容量、企业SSO、生产持久化、对象存储与外部系统连接器尚需按准入条件接入。")
    add_table(doc, ["待决事项", "需要确认", "责任方"], [
        ("正式GraphRAG", "Qwen建图/嵌入模型、独立Key、数据出域、预算、三审身份与索引制品库。", "华美业务、医务、合规、IT与项目组。"),
        ("试点边界", "院区、首发渠道、项目域、服务时段、人工兜底SLA与值班表。", "华美业务负责人。"),
        ("运营基线", "连续四周人工分钟、首响、知识查询、记录/跟进时间与服务质量数据。", "华美运营与项目经理。"),
        ("生产基础设施", "SSO、PostgreSQL、Redis、对象存储、监控、备份、域名与网络边界。", "华美IT、安全与SRE。"),
    ], [1700, 4750, 2910])
    add_para(doc, "—— 完 ——", color=MUTED, size=9.5)
    return doc


def technical_document():
    doc = Document()
    configure_document(doc, "华美一期售前智能体技术方案", "技术方案")
    add_cover(doc, "华美一期售前智能体\n技术方案", "独立服务、证据驱动检索与可控发布的技术实现方案", "技术架构、数据治理、部署运维与准入设计")

    add_section_title(doc, "1. 技术目标与设计原则", "本方案以“先形成独立可交付的售前智能体系统，再逐步扩展渠道和业务连接器”为技术路线。")
    add_table(doc, ["原则", "技术落点"], [
        ("服务独立", "售前智能体以独立服务对外提供统一会话协议，网页只是首个客户端；企业微信等渠道通过适配器接入。"),
        ("事实分路", "通用科普、企业知识、结构化业务事实、关系推理、媒体与受控写操作分别路由，避免单一路径泛化。"),
        ("证据优先", "企业事实必须绑定已发布知识、来源、版本、时效和权限；无法证实即降级、澄清或人工接管。"),
        ("发布不可变", "知识、索引、模型配置和评测报告形成可回溯发布单元；在线服务只读，不在对话中建图。"),
        ("安全失败关闭", "GraphRAG、模型、数据库、授权或连接器异常时不得虚构成功，切换至安全降级或人工路径。"),
        ("可运营可观测", "记录意图、路由、检索、证据、模型、耗时、成本、失败原因及Bad Case，支持持续评测。"),
    ], [1850, 7510])

    add_section_title(doc, "2. 总体架构")
    add_para(doc, "系统采用前后端分离与服务化架构。当前阶段不接入外部业务系统，但在领域服务与渠道层预留标准接口，避免网页实现与智能体能力耦合。")
    add_table(doc, ["架构层", "主要组件", "职责"], [
        ("交互层", "客户网页、管理后台", "展示对话、资料预览、运营工作台、权限内管理与质量反馈。"),
        ("接入层", "Web网关、未来渠道适配器", "统一身份、限流、请求标准化、媒体安全访问和渠道事件接入。"),
        ("智能体层", "售前智能体、意图/情绪、状态机、动作规划", "多轮理解、动态路由、回答组装、转人工与受控操作提案。"),
        ("知识层", "元数据RAG、结构化知识、GraphRAG、媒体服务", "提供企业事实、关系推理、资料引用、版本与时效校验。"),
        ("领域与协同层", "会话、线索、预约意向、排班、接管、Outbox", "管理业务状态、幂等、人工协同、审计及未来外部连接器。"),
        ("数据与平台层", "PostgreSQL、Redis、对象存储、监控、密钥管理", "持久化、缓存、任务、媒体、观测、备份和安全治理。"),
    ], [1650, 2750, 5210])
    add_process(doc, [
        ("客户请求", "网页或未来渠道\n统一协议"),
        ("安全上下文", "身份、租户、限流\n会话加载"),
        ("智能体编排", "意图、情绪、状态\n动态路由"),
        ("事实获取", "RAG/结构化/\nGraphRAG/媒体"),
        ("回答与动作", "证据过滤、短句输出\n接管或确认提案"),
        ("审计与指标", "路由、版本、耗时\n反馈与Bad Case"),
    ])

    add_section_title(doc, "3. 售前智能体运行时设计")
    add_subtitle(doc, "3.1 会话状态与待选择槽位")
    add_para(doc, "会话状态保存当前任务、话题版本、已确认槽位、待补字段、用户偏好、情绪信号和已完成步骤。用户的纠正、插话或简短选择先匹配待选择槽位，再进行通用意图分类。例如“自体脂肪方式”应写入丰胸方式槽位，而不是回退到“请说明部位”。")
    add_table(doc, ["状态字段", "示例", "处理规则"], [
        ("当前目标", "胸部整形咨询", "新主题达到切换阈值后创建新话题版本。"),
        ("已确认槽位", "方式=自体脂肪；偏好=自然", "后续回答必须继承，不重复询问。"),
        ("待选择槽位", "假体/自体脂肪二选一", "短语、序号、同义词先匹配候选项。"),
        ("澄清缺口", "预算或到院时间未提供", "只有确实影响下一步时才追问1—2项。"),
        ("情绪与风险", "焦虑、急迫、投诉、医疗高风险", "影响语气、转人工与安全策略，不改变事实结论。"),
    ], [1800, 3000, 4560])
    add_subtitle(doc, "3.2 路由矩阵")
    add_table(doc, ["问题类型", "检索/执行策略", "关键约束"], [
        ("稳定通用科普", "generic_model_education", "不伪造院方事实；不做个体诊断。"),
        ("医生、项目、产品、院区", "rag_metadata", "仅使用发布态企业知识及其来源。"),
        ("价格、活动、排班、名额、会员", "structured_business_data", "优先权威结构化源；未连接时提示需核验。"),
        ("跨医生/项目比较、多跳关系", "graph_rag", "仅在实体关系/关联推理需要时调用；失败自动降级。"),
        ("公开资料展示", "knowledge_media", "按资料权限签发短时预览链接。"),
        ("预约、报名、工单", "confirmation_required", "确认、幂等、审计、可取消窗口和连接器缺一不可。"),
        ("投诉、医疗风险、隐私、无法回答", "human_handoff", "携带摘要、证据、风险和已收集字段；AI转静默。"),
    ], [1850, 3200, 4310])

    add_section_title(doc, "4. 企业知识、RAG 与 GraphRAG 协同")
    add_table(doc, ["层级", "存放内容", "适用问题", "更新方式"], [
        ("缓存", "短时热点答案、检索结果、会话状态", "重复查询与并发去重", "TTL、版本与租户作为缓存键；发布后自动失效。"),
        ("结构化数据库", "产品字段、价格、活动、排班、审批、预约、角色", "时效性强、可过滤、需精确核验的信息", "事务更新、审批与生效时间控制。"),
        ("知识库/RAG", "医生、项目、产品公开资料、服务流程、PDF/图片解析文本", "单实体事实、资料定位、公开引用", "文档版本、元数据、审核发布包。"),
        ("GraphRAG", "医生—项目—术式—产品—院区等关系、社区摘要与关系证据", "多实体比较、多跳关系、关联发现", "从已批准发布包离线构建不可变索引。"),
    ], [1500, 3150, 2600, 2110], font_size=8.9)
    add_callout(doc, "GraphRAG 使用原则", "GraphRAG不是默认检索方案。只有实体关系、多跳推理和项目关系发现走图谱；价格、活动名额、会员状态等优先结构化数据库与元数据过滤。", "brand")
    add_subtitle(doc, "4.1 GraphRAG 建图与发布")
    add_numbered(doc, [
        "从审核通过的对客候选资料生成唯一、不可变的公开安全知识包；客户原话、联系方式、病史、案例素材与内部运营汇总均不进入建图输入。",
        "离线执行实体/关系抽取、社区发现、向量化和索引生成；使用独立Qwen建图与嵌入凭据，记录模型、端点摘要、Token与成本。",
        "对候选索引进行完整性校验、来源一致性校验、回归评测、对抗安全测试和20并发验证。",
        "业务、医务、合规三审回执与索引Manifest绑定，由索引管理员绑定、发布人激活；线上实例只读加载新版本。",
        "切换后保留上一版本；异常时按发布记录回滚，不对旧索引做原地改写。",
    ])
    add_para(doc, "当前开发候选目录为925条公开安全记录，预构建索引包含464个实体和2,812条关系；其属于开发基线，不等同于已获三审批准的正式索引。")

    add_section_title(doc, "5. 模型与多模态策略")
    add_table(doc, ["任务", "推荐模型策略", "使用边界"], [
        ("在线售前对话", "以平衡模型承担普通生成；复杂多项目比较或高价值资料解析可路由至更强模型。", "企业事实仍由证据约束，模型不得替代知识发布。"),
        ("GraphRAG建图", "候选使用 Qwen 3.8 Max 进行实体/关系抽取；首次正式建图前做质量/成本评测。", "离线任务；需数据出域与预算授权。"),
        ("文本向量", "使用专用Embedding模型，如 qwen3.7-text-embedding 或经评测确认的 text-embedding-v4。", "更换Embedding必须全量重嵌入并新建索引版本，禁止混用。"),
        ("图片/PDF", "图片型资料先隔离解析；必要时使用视觉理解模型提取结构化信息，并将公开媒体存入对象存储。", "只处理已批准公开资料；原始受限文件不对客开放。"),
    ], [1750, 4300, 3310])
    add_callout(doc, "模型升级", "Qwen 3.8的发布不触发自动索引替换。模型、提示词、检索算法或知识库任一变更，必须构建候选版本并跑固定回归集、真实问题集、对抗安全集、长对话集、工具执行集与人工盲评。", "caution")

    add_section_title(doc, "6. 数据模型、存储与媒体服务")
    add_table(doc, ["数据域", "主存储", "关键保护"], [
        ("会话与对话状态", "PostgreSQL + Redis", "租户隔离、最小化保存、会话级访问控制、日志脱敏。"),
        ("联系方式与授权", "独立加密保险箱", "独立Keyring、显式查看、双角色权限、实时授权复核、追加审计。"),
        ("知识与审批", "PostgreSQL + 私有对象存储", "版本、来源、有效期、三审、Hash、审计与只读发布包。"),
        ("图片/PDF公开资料", "OSS/对象存储 + 媒体索引", "病毒扫描、ACL、短时签名URL、内容指纹与对客范围控制。"),
        ("索引制品", "版本化、不可覆盖对象存储", "Manifest、SHA-256、只读挂载、索引版本与知识Release绑定。"),
        ("指标与审计", "Prometheus/日志/审计存储", "不写入敏感明文；跨副本聚合、留存与访问控制。"),
    ], [1900, 3000, 4460])

    add_section_title(doc, "7. 安全、权限与审计")
    add_bullets(doc, [
        "生产身份采用企业OIDC/SSO，后端基于实名主体、tenant、院区、科室和角色做强制授权；前端不拥有扩大权限的能力。",
        "所有写操作使用幂等键、服务端权限校验和审计事件；外部业务写入通过Outbox异步投递并记录回执、重试与人工对账。",
        "模型调用只允许到已批准的HTTPS出口；建图和在线生成使用独立API Key，密钥经Secrets Manager注入，不进入代码、镜像或日志。",
        "防护提示词注入、越权检索、跨租户缓存污染、未授权媒体访问和敏感信息外泄；异常默认降级。",
        "客户可在授权范围内查询、更正、删除其长期记忆和资料；非授权的客户资料不得进入通用向量检索或训练样本。",
    ])
    add_table(doc, ["角色", "技术权限边界"], [
        ("内容专员/科室主管", "仅维护授权科室的草稿与标签；不得直接对客发布。"),
        ("业务/医务/合规审核人", "分别完成独立实名审核；不得互相代审。"),
        ("索引管理员", "绑定已批准的GraphRAG索引；不得修改原文或审核结论。"),
        ("发布人", "只激活已完成三审、索引绑定和评测通过的发布包。"),
        ("人工坐席", "认领后按最小权限处理会话；联系方式查看需额外授权并全审计。"),
    ], [2750, 6610])

    add_section_title(doc, "8. 部署、环境与容量设计")
    add_table(doc, ["环境", "目的", "关键配置"], [
        ("开发环境", "功能开发、模拟数据、模型与建图候选验证。", "单机Docker可用；独立开发业务空间与密钥；不承载真实客户数据。"),
        ("测试/预生产", "真实模型、正式候选知识、20并发30分钟、故障注入、恢复和灰度验证。", "PostgreSQL、Redis、对象存储、监控、受控模型出口、OIDC与备份恢复。"),
        ("生产环境", "低风险试点与后续稳定运营。", "多副本服务、只读GraphRAG索引、HA数据底座、告警、审计、灰度与回滚。"),
    ], [1650, 3300, 4410])
    add_table(doc, ["指标", "目标/门槛", "说明"], [
        ("并发基线", "20名用户并发，稳态30分钟", "真实模型、正式知识版本与完整监控条件下验证。"),
        ("时延", "简单问答首字P95≤2秒；GraphRAG P95≤2秒；完整链路P95≤10秒", "超时进入已审核RAG/模板/人工降级，不阻塞会话。"),
        ("可用性", "生产目标≥99.9%", "需以实际SLA、监控、故障演练和依赖等级确认。"),
        ("容量架构", "Web/API、智能体、GraphRAG均可水平扩展", "共享状态进入PostgreSQL/Redis；索引以只读制品挂载；缓存按租户和版本隔离。"),
    ], [1650, 3400, 4310])

    add_section_title(doc, "9. 可观测性、评测与持续交付")
    add_table(doc, ["对象", "采集内容", "运营用途"], [
        ("智能体", "意图、槽位、路由、模型、耗时、成本、降级、转人工、回答反馈。", "发现对话偏差与Bad Case，支持问题复盘。"),
        ("检索与GraphRAG", "索引版本、召回、证据准入、缓存、净化查询SHA、错误率与P95。", "识别索引故障、知识缺口和发布回归。"),
        ("业务工作台", "认领、处理、预约核验、工时、SLA与回执。", "验证人工投入下降和服务质量。"),
        ("安全与数据", "授权拒绝、敏感拦截、媒体访问、密钥/依赖状态、审计完整性。", "支持安全告警、取证与合规复核。"),
    ], [1800, 4500, 3060])
    add_subtitle(doc, "9.1 变更门禁")
    add_bullets(doc, [
        "代码变更：单元、接口、E2E、权限、隐私、安全和回归测试通过后方可进入测试环境。",
        "提示词、模型、检索算法或知识变更：必须运行固定回归集、线上真实问题集、对抗安全集、长对话集、工具执行集和人工盲评。",
        "GraphRAG索引变更：必须验证知识包SHA、批准记录SHA、Manifest、证据一致性、容量和回滚路径。",
        "生产发布：需变更记录、监控验证、备份恢复、人工接管值班和项目委员会Go/Conditional Go。",
    ])

    add_section_title(doc, "10. 正式 GraphRAG 准入与运行流程")
    add_process(doc, [
        ("准备", "批准知识包\n数据出域/预算"),
        ("离线建图", "Qwen抽取、嵌入\n生成候选制品"),
        ("质量验证", "完整性、回归、安全\n20并发验证"),
        ("三审绑定", "业务、医务、合规\n不可变回执"),
        ("灰度激活", "索引管理员绑定\n发布人切换"),
        ("观测回滚", "指标监控、问题\n可退回旧版本"),
    ])
    add_table(doc, ["必备输入", "说明"], [
        ("Qwen模型配置", "建图/查询模型、Embedding模型、Endpoint、独立API Key、RPM/TPM/并发额度。"),
        ("数据出域授权", "数据地域、接收方、留存、训练使用、删除条款、允许字段范围及审批编号。"),
        ("成本控制", "构建预算上限、Token阈值、阿里云费用预算、50%/80%/100%告警联系人和停止机制。"),
        ("知识治理", "业务、医务、合规三名不同实名审核人、知识包SHA、有效期、发布与回滚版本。"),
        ("基础设施", "不可变索引制品库、Secrets Manager、受控模型出口、监控、备份与测试环境。"),
    ], [2050, 7310])

    add_section_title(doc, "11. 当前边界、风险与实施建议")
    add_table(doc, ["项目项", "当前状态", "实施建议"], [
        ("正式GraphRAG", "开发候选索引存在；正式LLM索引未准入。", "完成Qwen配置、数据出域、预算与三审后，在开发/测试环境构建候选索引。"),
        ("外部业务写入", "当前为内部预约与Outbox框架；未接CRM/排班。", "保持关闭；后续以测试租户、字段映射、幂等、回执和人工确认先做沙箱联调。"),
        ("生产数据底座", "本地可运行；正式HA与客户环境演练未完成。", "测试环境先部署PostgreSQL、Redis、对象存储、监控与备份，完成故障演练。"),
        ("模型升级", "Qwen 3.8可作为建图抽取候选，Embedding无需随版本自动替换。", "新模型构建独立候选索引，与基线比对质量、成本和延迟后再发布。"),
        ("多渠道", "网页为首发渠道。", "智能体API保持渠道无关；企业微信等通过适配器和统一会话协议后续接入。"),
    ], [1700, 3400, 4260])
    add_callout(doc, "上线结论", "在正式知识Release、GraphRAG索引、真实模型容量、身份权限、数据底座、监控告警、备份恢复和人工接管演练完成前，系统仅应作为开发/受控外测环境，不得被描述为正式对客生产服务。", "caution")
    add_para(doc, "—— 完 ——", color=MUTED, size=9.5)
    return doc


def main():
    product_path = OUT / "华美一期售前智能体产品设计方案_V1.0.docx"
    technical_path = OUT / "华美一期售前智能体技术方案_V1.0.docx"
    product_document().save(product_path)
    technical_document().save(technical_path)
    print(product_path)
    print(technical_path)


if __name__ == "__main__":
    main()
